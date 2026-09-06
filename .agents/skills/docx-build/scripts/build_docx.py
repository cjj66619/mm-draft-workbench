#!/usr/bin/env python3
"""build_docx.py — 把 paper/sections/*.md（Markdown + LaTeX 公式）合成一份 Word 初稿 paper/main.docx。

设计原则（mm-draft-workbench）：
- 初稿阶段：Markdown 章节是源，随时可重生成；
- 生成后 `--freeze`：Word 成为唯一真源，之后只改 Word，本脚本拒绝再覆盖（除非 --force）。
  冻结状态记录在 paper/DOCX_FREEZE.json，后续人/智能体先看它再动手。

流程：
1. 读取 paper/paper.yaml（题目、关键词、章节顺序）与 paper/sections/*.md（按文件名排序，或 yaml 指定）。
   摘要放 paper/sections/00_abstract.md（不写标题，空行分段）。
2. 解析交叉引用：图 `![图题](../../figures/<id>/<id>.pdf){#fig:xxx}`、表 `Table: 表题 {#tbl:xxx}`、
   公式 `$$ ... $$ {#eq:xxx}`；正文里 `@fig:xxx` / `@tbl:xxx` / `@eq:xxx` → 图N / 表N / 式(N)。
3. 正文里引用的 PDF 图转成同名 PNG（Word 不支持 PDF 图；PNG 写入 dpi 保持物理尺寸）。
4. pandoc（markdown → docx），公式输出为可编辑 OMML（Word 原生公式）。
5. python-docx 套用中文论文正文格式：A4 四边 2.5 cm、宋体/Times New Roman 12 pt 首行缩进 2 字符、
   一级标题“一、”居中黑体 14 pt、二/三级 “1.1”/“1.1.1” 12 pt、图题下/表题上 11 pt 加粗、三线表、
   公式居中编号右对齐、摘要页、目录页（TOC 域 + 静态缓存）、页脚居中 9 pt 页码。
   有当届官方 Word 模板时 `--reference 官方模板.docx` 直接套用其页面设置。
6. 有 soffice/LibreOffice 时 `--pdf` 顺手转一份 PDF 用于渲染抽检、统计页数并回填目录页码。
7. 审计：模板占位符、内部文件名泄露（reports/、figures/、AGENTS.md…）、页面设置、图表题数量。

用法（项目根目录，Windows/Linux/macOS 同）：
    python tools/build_docx.py                       # 生成 paper/main.docx
    python tools/build_docx.py --pdf --strict        # + LibreOffice 渲染抽检 + 严格审计
    python tools/build_docx.py --freeze              # 生成并冻结（之后 Word 为唯一真源）
    python tools/build_docx.py --status              # 查看冻结状态
    python tools/build_docx.py --force               # 冻结后仍要重生成（会覆盖人工 Word 修改！先备份）
    python tools/build_docx.py --unfreeze            # 解除冻结（明确回到 Markdown 为源）

依赖：pandoc（必需）、python-docx（必需）、pymupdf 或 pdftoppm（PDF 图转 PNG）、soffice（可选）。
Windows 上 pandoc 装在默认位置即可被探测到，也可用环境变量 PANDOC_BIN / SOFFICE_BIN 指定。
"""

from __future__ import annotations

import argparse
import copy
import datetime as _dt
import hashlib
import json
import os
import platform
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

try:
    import docx
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Twips
except ImportError:  # pragma: no cover
    sys.exit("缺少 python-docx：pip install python-docx")

# --------------------------------------------------------------------------- 常量（中文论文正文格式）
PAGE_W, PAGE_H = 11906, 16838          # A4，twips
MARGIN = 1418                          # 2.5 cm
HEADER_DIST, FOOTER_DIST = 851, 992
TEXT_W = PAGE_W - 2 * MARGIN           # 9070 版心宽
SONG, HEI, LI = "宋体", "黑体", "隶书"
TNR = "Times New Roman"
CODE_FONT = "Courier New"

CN_DIGITS = "零一二三四五六七八九"
FREEZE_FILE = "DOCX_FREEZE.json"
ABSTRACT_NAMES = ("00_abstract.md", "abstract.md", "摘要.md")

PDF2PNG_LUA = r"""
-- 把 image 的 .pdf 源改为同名 .png（由 build_docx.py 预先生成）
function Image(el)
  if el.src:match("%.pdf$") or el.src:match("%.PDF$") then
    el.src = el.src:gsub("%.[pP][dD][fF]$", ".png")
  end
  return el
end

local PAGE_BREAK = pandoc.RawBlock("openxml", '<w:p><w:r><w:br w:type="page"/></w:r></w:p>')

-- typst #pagebreak() -> Div.page-break；latex \newpage / \clearpage -> RawBlock；本脚本生成的摘要页 -> Div.pagebreak
function Div(el)
  if el.classes:includes("page-break") or el.classes:includes("pagebreak") then
    return PAGE_BREAK
  end
end

function RawBlock(el)
  if el.format == "latex" and (el.text:match("\\newpage") or el.text:match("\\clearpage")) then
    return PAGE_BREAK
  end
end

-- 图/表题自动编号：官方格式“图N xxx” / “表N xxx”（编号紧跟“图/表”，其后一个空格）
local fig_n, tab_n = 0, 0

local function prefix_caption(cap, label)
  if cap == nil or cap.long == nil or #cap.long == 0 then
    return cap
  end
  local first = cap.long[1]
  if first.t == "Plain" or first.t == "Para" then
    first.content:insert(1, pandoc.Str(label))
    first.content:insert(2, pandoc.Space())
  end
  return cap
end

function Figure(el)
  fig_n = fig_n + 1
  el.caption = prefix_caption(el.caption, "图" .. fig_n)
  return el
end

function Table(el)
  tab_n = tab_n + 1
  el.caption = prefix_caption(el.caption, "表" .. tab_n)
  return el
end

-- 参考文献 / 附录 一级标题不编号（Typst 的 numbering: none 不会传给 pandoc）
function Header(el)
  local txt = pandoc.utils.stringify(el)
  if el.level == 1 and (txt:match("^参考文献") or txt:match("^附录")) then
    el.classes:insert("unnumbered")
  end
  return el
end

-- 行间公式全篇连续编号；LaTeX \eqref{eq:x} / \ref{eq:x} 解析为 (N)
local eq_labels, eq_n = {}, 0

local function collect_math(el)
  if el.mathtype == "DisplayMath" then
    eq_n = eq_n + 1
    for lab in el.text:gmatch("\\label{([^}]+)}") do
      eq_labels[lab] = eq_n
    end
  end
  return nil
end

local function rewrite_ref(el)
  local ref = el.attributes["reference"]
  if ref and eq_labels[ref] then
    return pandoc.Str("(" .. eq_labels[ref] .. ")")
  end
end

-- 只含空 Span 的段落（Typst <label> 残留）删除
local function drop_empty(el)
  if #el.content == 0 then
    return {}
  end
  for _, x in ipairs(el.content) do
    if not (x.t == "Span" and #x.content == 0) then
      return nil
    end
  end
  return {}
end

return {
  { Math = collect_math },
  { Image = Image, Div = Div, RawBlock = RawBlock, Figure = Figure, Table = Table,
    Header = Header, Link = rewrite_ref, Para = drop_empty, Plain = drop_empty },
}
"""

TITLE_MARK = "\u200b\u200bTITLE\u200b\u200b"
ABSTRACT_MARK = "\u200b\u200bABSTRACT\u200b\u200b"
KEYWORDS_MARK = "\u200b\u200bKEYWORDS\u200b\u200b"

IMAGE_MD = re.compile(r'!\[[^\]]*\]\(([^)\s]+\.pdf)(?:\s+"[^"]*")?\)', re.I)
PLACEHOLDER = re.compile(r'\[(论文标题|学校名称|参赛队号|成员 ?[A-C]|关键词\d?|中文摘要内容[^\]]*|封面[^\]]*|TODO[^\]]*|待补[^\]]*)\]')
INTERNAL_NAMES = ("reports/", "figures/", "results/", "RESULTS_REPORT", "ANALYSIS_MODELING_REPORT", "CLAUDE.md", "AGENTS.md",
                  "HANDOFF.md", "plan.md", "todo.md", "make_figure.py", "review.json", "manifest.json")
LEAD_PHRASE = re.compile(r'^(针对问题[一二三四五六七八九十\d]+[，,：:])')

# 交叉引用：定义与引用
FIG_DEF = re.compile(r'!\[[^\]]*\]\([^)]*\)\s*\{[^}]*#(fig:[\w:.-]+)[^}]*\}')
TBL_DEF = re.compile(r'^(?:Table|表)\s*:\s*[^\n]*?\{[^}]*#(tbl:[\w:.-]+)[^}]*\}\s*$', re.M)
EQ_DEF = re.compile(r'\$\$(?:[^$]|\$(?!\$))*\$\$\s*(\{[^}]*#(eq:[\w:.-]+)[^}]*\})?')
REF = re.compile(r'(?<![\w@])@((?:fig|tbl|eq):[\w:.-]*[\w])')
ATTR_TAIL = re.compile(r'\s*\{[^}]*#(?:fig|tbl|eq):[\w:.-]+[^}]*\}')


@dataclass
class FrontMatter:
    title: str = ""
    abstract: list[str] = field(default_factory=list)
    keywords: list[str] = field(default_factory=list)


def run(cmd: list[str], **kw) -> subprocess.CompletedProcess:
    return subprocess.run(cmd, text=True, capture_output=True, encoding="utf-8", errors="replace", **kw)


def cn_number(n: int) -> str:
    """1 -> 一, 10 -> 十, 12 -> 十二, 21 -> 二十一（一级标题编号用）。"""
    if n <= 0 or n >= 100:
        return str(n)
    if n < 10:
        return CN_DIGITS[n]
    tens, ones = divmod(n, 10)
    return ("" if tens == 1 else CN_DIGITS[tens]) + "十" + (CN_DIGITS[ones] if ones else "")


# --------------------------------------------------------------------------- 外部程序探测（跨平台）
def find_exe(name: str, env_var: str, win_candidates: list[str] = (), mac_candidates: list[str] = ()) -> str | None:
    env = os.environ.get(env_var)
    if env and Path(env).exists():
        return env
    hit = shutil.which(name)
    if hit:
        return hit
    cands = win_candidates if platform.system() == "Windows" else (mac_candidates if platform.system() == "Darwin" else [])
    for c in cands:
        p = Path(os.path.expandvars(c))
        if p.exists():
            return str(p)
    return None


def find_pandoc() -> str | None:
    return find_exe("pandoc", "PANDOC_BIN",
                    [r"%ProgramFiles%\Pandoc\pandoc.exe", r"%LocalAppData%\Pandoc\pandoc.exe"],
                    ["/opt/homebrew/bin/pandoc", "/usr/local/bin/pandoc"])


def find_soffice() -> str | None:
    return find_exe("soffice", "SOFFICE_BIN",
                    [r"%ProgramFiles%\LibreOffice\program\soffice.exe", r"%ProgramFiles(x86)%\LibreOffice\program\soffice.exe"],
                    ["/Applications/LibreOffice.app/Contents/MacOS/soffice"]) or shutil.which("libreoffice")


# --------------------------------------------------------------------------- 源码收集（Markdown）
def load_meta(paper: Path) -> dict:
    """paper/paper.yaml：title / keywords / sections / reference。没有 PyYAML 时用最小解析（键: 值、- 列表）。"""
    f = paper / "paper.yaml"
    if not f.exists():
        return {}
    text = f.read_text(encoding="utf-8")
    try:
        import yaml  # type: ignore
        data = yaml.safe_load(text) or {}
        return data if isinstance(data, dict) else {}
    except ImportError:
        pass
    data: dict = {}
    key = None
    for line in text.splitlines():
        if not line.strip() or line.lstrip().startswith("#"):
            continue
        if line.startswith((" ", "\t", "-")) and key:
            item = line.strip().lstrip("-").strip().strip('"\'')
            data.setdefault(key, [])
            if isinstance(data[key], list):
                data[key].append(item)
            continue
        if ":" in line:
            key, _, val = line.partition(":")
            key, val = key.strip(), val.strip().strip('"\'')
            data[key] = val if val else []
    return data


def collect_sections(paper: Path, meta: dict) -> tuple[Path | None, list[Path]]:
    """返回 (摘要文件, 正文章节列表)。顺序：paper.yaml 的 sections，否则按文件名排序。"""
    sec_dir = paper / "sections"
    if not sec_dir.is_dir():
        sys.exit(f"{sec_dir} 不存在：请把章节写成 paper/sections/NN_xxx.md")
    order = meta.get("sections")
    if isinstance(order, list) and order:
        files = []
        for name in order:
            p = sec_dir / str(name)
            if not p.suffix:
                p = p.with_suffix(".md")
            if p.exists():
                files.append(p)
            else:
                print(f"[warn] paper.yaml sections 中的 {name} 不存在，跳过")
    else:
        files = sorted(p for p in sec_dir.glob("*.md") if not p.name.startswith("_"))
    abstract = next((p for p in files if p.name in ABSTRACT_NAMES), None)
    body = [p for p in files if p is not abstract]
    return abstract, body


def extract_front_matter(paper: Path, meta: dict, abstract_file: Path | None) -> FrontMatter:
    fm = FrontMatter()
    fm.title = str(meta.get("title") or "").strip()
    kws = meta.get("keywords") or []
    if isinstance(kws, str):
        kws = re.split(r"[;；,，]", kws)
    fm.keywords = [str(k).strip() for k in kws if str(k).strip()]
    if abstract_file is not None:
        text = abstract_file.read_text(encoding="utf-8")
        text = re.sub(r"^---\n.*?\n---\n", "", text, flags=re.S)          # 去 YAML 头
        text = re.sub(r"^#+\s*摘\s*要[：:]?\s*$", "", text, flags=re.M)   # 去“# 摘要”标题
        m = re.search(r"^\**关键词[：:]\**\s*(.+)$", text, flags=re.M)
        if m and not fm.keywords:
            fm.keywords = [k.strip() for k in re.split(r"[;；,，]|\u3000+|\s{2,}", m.group(1)) if k.strip()]
        if m:
            text = text[:m.start()] + text[m.end():]
        fm.abstract = [p.strip() for p in re.split(r"\n\s*\n", text.strip()) if p.strip()]
    return fm


def resolve_crossrefs(sections: list[Path]) -> tuple[dict[str, str], list[str]]:
    """按出现顺序给 fig/tbl/eq 标签编号（与 pandoc 图表编号、number_equations 公式编号一致）。

    返回 ({label: '图3'|'表2'|'式(5)'}, [改写后章节文本])。未定义的 @ref 原样保留并告警。
    """
    numbers: dict[str, str] = {}
    fig_n = tbl_n = eq_n = 0
    stripped: list[str] = []
    for sec in sections:
        text = sec.read_text(encoding="utf-8")
        code_free = re.sub(r"```.*?```", lambda m: " " * len(m.group(0)), text, flags=re.S)
        events: list[tuple[int, str, str | None]] = []
        for m in FIG_DEF.finditer(code_free):
            events.append((m.start(), "fig", m.group(1)))
        for m in TBL_DEF.finditer(code_free):
            events.append((m.start(), "tbl", m.group(1)))
        for m in EQ_DEF.finditer(code_free):
            events.append((m.start(), "eq", m.group(2)))
        for _, kind, label in sorted(events):
            if kind == "fig":
                fig_n += 1
                numbers[label] = f"图{fig_n}"
            elif kind == "tbl":
                tbl_n += 1
                numbers[label] = f"表{tbl_n}"
            else:
                eq_n += 1
                if label:
                    numbers[label] = f"式({eq_n})"
        stripped.append(text)
    out: list[str] = []
    for text in stripped:
        def _sub(m: re.Match) -> str:
            lab = m.group(1)
            if lab in numbers:
                return numbers[lab]
            print(f"[warn] 未定义的交叉引用 @{lab}，原样保留")
            return m.group(0)
        text = REF.sub(_sub, text)
        text = ATTR_TAIL.sub("", text)
        out.append(text)
    return numbers, out


def convert_pdf_figures(texts: list[str], base: Path, dpi: int = 300) -> list[Path]:
    """把章节里引用的 PDF 图转成 PNG（同目录同名，写入 dpi 元数据以保持物理尺寸）。路径相对 paper/sections/。"""
    try:
        import pymupdf  # type: ignore
    except ImportError:
        try:
            import fitz as pymupdf  # type: ignore
        except ImportError:
            pymupdf = None
    pdftoppm = shutil.which("pdftoppm")
    made: list[Path] = []
    seen: set[Path] = set()
    for text in texts:
        for m in IMAGE_MD.finditer(text):
            pdf = (base / m.group(1)).resolve()
            if pdf in seen:
                continue
            seen.add(pdf)
            png = pdf.with_suffix(".png")
            if not pdf.exists():
                print(f"[warn] 图片不存在: {pdf}")
                continue
            if png.exists() and png.stat().st_mtime >= pdf.stat().st_mtime:
                continue
            if pymupdf is not None:
                doc = pymupdf.open(pdf)
                pix = doc[0].get_pixmap(dpi=dpi)
                pix.set_dpi(dpi, dpi)
                pix.save(png)
                doc.close()
            elif pdftoppm:
                run([pdftoppm, "-png", "-r", str(dpi), "-singlefile", str(pdf), str(png.with_suffix(""))])
            else:
                print(f"[warn] 无 pymupdf/pdftoppm，无法转换 {pdf}（pip install pymupdf）")
                continue
            made.append(png)
    return made


def front_matter_source(fm: FrontMatter) -> str:
    """标题 / 摘要 / 关键词写成 Markdown 交给 pandoc 一起转，摘要里的公式也能保留。

    三个零宽标记在 style_front_matter() 中被识别并替换成摘要页格式。
    """
    title = fm.title or "[论文标题]"
    abstract = fm.abstract or ["[中文摘要内容]"]
    keywords = fm.keywords or ["[关键词]"]
    return (
        f"{TITLE_MARK}{title}\n\n"
        f"{ABSTRACT_MARK}摘\u3000要：\n\n"
        + "\n\n".join(abstract)
        + f"\n\n{KEYWORDS_MARK}关键词：" + "\u3000\u3000".join(keywords)
        + "\n\n::: {.page-break}\n:::\n\n"
    )


# --------------------------------------------------------------------------- 冻结机制
def sha256_of(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def sections_digest(files: list[Path]) -> str:
    h = hashlib.sha256()
    for f in files:
        h.update(f.name.encode("utf-8"))
        h.update(f.read_bytes())
    return h.hexdigest()[:16]


def load_freeze(paper: Path) -> dict:
    f = paper / FREEZE_FILE
    if not f.exists():
        return {}
    try:
        return json.loads(f.read_text(encoding="utf-8"))
    except json.JSONDecodeError:
        return {}


def save_freeze(paper: Path, data: dict) -> None:
    (paper / FREEZE_FILE).write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def freeze_status(paper: Path, out: Path) -> tuple[str, dict]:
    """返回 (状态, 记录)。状态：none | draft | frozen | frozen-edited | frozen-missing。"""
    rec = load_freeze(paper)
    if not rec:
        return "none", rec
    if not rec.get("frozen"):
        return "draft", rec
    if not out.exists():
        return "frozen-missing", rec
    return ("frozen-edited" if sha256_of(out) != rec.get("docx_sha256") else "frozen"), rec


def print_status(paper: Path, out: Path) -> None:
    st, rec = freeze_status(paper, out)
    msg = {
        "none": "未生成过 DOCX（Markdown 为源，可直接 build）",
        "draft": f"草稿态：DOCX 于 {rec.get('generated_at')} 由 Markdown 生成，可重生成；确认后用 --freeze 冻结",
        "frozen": f"已冻结（{rec.get('frozen_at')}）：Word 为唯一真源，请直接编辑 {out.name}；不要重生成",
        "frozen-edited": f"已冻结且 Word 已被人工修改（哈希不同于 {rec.get('frozen_at')} 冻结时）：只能改 Word；--force 会丢弃这些修改",
        "frozen-missing": f"已冻结但 {out.name} 不存在：请从备份恢复，或 --unfreeze 后重生成",
    }[st]
    print(f"[status] {st}: {msg}")
    if rec.get("sections_digest"):
        print(f"[status] 冻结时章节摘要 {rec['sections_digest']}（Markdown 之后的改动不会进入 Word）")


# --------------------------------------------------------------------------- OOXML 小工具
# Word 对子元素顺序敏感（乱序会报“无法读取的内容”），按 ECMA-376 schema 顺序插入。
_ORDER: dict[str, list[str]] = {
    "w:pPr": ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl", "numPr",
              "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap",
              "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd",
              "snapToGrid", "spacing", "ind", "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc",
              "textDirection", "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr",
              "sectPr", "pPrChange"],
    "w:rPr": ["rStyle", "rFonts", "b", "bCs", "i", "iCs", "caps", "smallCaps", "strike", "dstrike", "outline",
              "shadow", "emboss", "imprint", "noProof", "snapToGrid", "vanish", "webHidden", "color", "spacing",
              "w", "kern", "position", "sz", "szCs", "highlight", "u", "effect", "bdr", "shd", "fitText",
              "vertAlign", "rtl", "cs", "em", "lang", "eastAsianLayout", "specVanish", "oMath"],
    "w:tblPr": ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize", "tblStyleColBandSize",
                "tblW", "jc", "tblCellSpacing", "tblInd", "tblBorders", "shd", "tblLayout", "tblCellMar",
                "tblLook", "tblCaption", "tblDescription"],
    "w:tcPr": ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap", "tcMar",
               "textDirection", "tcFitText", "vAlign", "hideMark"],
    "w:tblBorders": ["top", "start", "left", "bottom", "end", "right", "insideH", "insideV"],
    "w:tcBorders": ["top", "start", "left", "bottom", "end", "right", "insideH", "insideV", "tl2br", "tr2bl"],
    "w:pBdr": ["top", "left", "bottom", "right", "between", "bar"],
    "w:tblCellMar": ["top", "start", "left", "bottom", "end", "right"],
    "w:sectPr": ["headerReference", "footerReference", "footnotePr", "endnotePr", "type", "pgSz", "pgMar",
                 "paperSrc", "pgBorders", "lnNumType", "pgNumType", "cols", "formProt", "vAlign", "noEndnote",
                 "titlePg", "textDirection", "bidi", "rtlGutter", "docGrid"],
    "w:style": ["name", "aliases", "basedOn", "next", "link", "autoRedefine", "hidden", "uiPriority", "semiHidden",
                "unhideWhenUsed", "qFormat", "locked", "personal", "personalCompose", "personalReply", "rsid",
                "pPr", "rPr", "tblPr", "trPr", "tcPr", "tblStylePr"],
    "w:p": ["pPr"],
    "w:tc": ["tcPr"],
    "w:tbl": ["tblPr", "tblGrid", "tr"],
    "w:rPrDefault": ["rPr"],
}


def _local(tag: str) -> str:
    return tag.rsplit("}", 1)[-1] if "}" in tag else tag.split(":")[-1]


def _insert_ordered(parent, el) -> None:
    order = _ORDER.get(f"w:{_local(parent.tag)}")
    name = _local(el.tag)
    if order is None or name not in order:
        parent.append(el)
        return
    rank = order.index(name)
    for sib in parent:
        sib_name = _local(sib.tag)
        if sib_name in order and order.index(sib_name) > rank:
            sib.addprevious(el)
            return
    parent.append(el)


def _child(parent, tag: str, **attrs):
    """取或按 schema 顺序新建子元素（属性用 w: 前缀名，如 val="x" -> w:val）。"""
    el = parent.find(qn(tag))
    if el is None:
        el = OxmlElement(tag)
        _insert_ordered(parent, el)
    for k, v in attrs.items():
        el.set(qn(f"w:{k}"), str(v))
    return el


def _set_fonts(rpr, west: str = TNR, east: str = SONG) -> None:
    rf = _child(rpr, "w:rFonts")
    for a in ("ascii", "hAnsi", "cs"):
        rf.set(qn(f"w:{a}"), west)
    rf.set(qn("w:eastAsia"), east)
    for a in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        rf.attrib.pop(qn(f"w:{a}"), None)


def _set_size(rpr, half_pts: int) -> None:
    _child(rpr, "w:sz", val=half_pts)
    _child(rpr, "w:szCs", val=half_pts)


def _set_bold(rpr, bold: bool) -> None:
    for tag in ("w:b", "w:bCs"):
        el = rpr.find(qn(tag))
        if bold:
            if el is None:
                el = OxmlElement(tag)
                _insert_ordered(rpr, el)
            el.attrib.pop(qn("w:val"), None)
        elif el is not None:
            rpr.remove(el)


def _clear(rpr_or_ppr, *tags: str) -> None:
    for t in tags:
        for el in rpr_or_ppr.findall(qn(t)):
            rpr_or_ppr.remove(el)


def _ppr(el):
    return _child(el, "w:pPr")


def _para_props(ppr, *, jc: str | None = None, first_line_chars: int | None = None,
                left: int | None = None, hanging: int | None = None,
                before: int | None = None, after: int | None = None,
                line: int | None = None, line_rule: str | None = None,
                keep_next: bool | None = None, keep_lines: bool | None = None,
                outline: int | None = None) -> None:
    if keep_next:
        _child(ppr, "w:keepNext")
    if keep_lines:
        _child(ppr, "w:keepLines")
    if before is not None or after is not None or line is not None:
        sp = _child(ppr, "w:spacing")
        if before is not None:
            sp.set(qn("w:before"), str(before))
            sp.attrib.pop(qn("w:beforeLines"), None)
        if after is not None:
            sp.set(qn("w:after"), str(after))
            sp.attrib.pop(qn("w:afterLines"), None)
        if line is not None:
            sp.set(qn("w:line"), str(line))
            sp.set(qn("w:lineRule"), line_rule or "auto")
    if first_line_chars is not None or left is not None or hanging is not None:
        ind = _child(ppr, "w:ind")
        for a in ("firstLine", "firstLineChars", "left", "leftChars", "hanging", "hangingChars", "start", "startChars"):
            ind.attrib.pop(qn(f"w:{a}"), None)
        if first_line_chars is not None:
            ind.set(qn("w:firstLineChars"), str(first_line_chars))
            ind.set(qn("w:firstLine"), str(first_line_chars * 12 // 5))  # 12 pt 字号下 100 chars = 240 twips
        if left is not None:
            ind.set(qn("w:left"), str(left))
        if hanging is not None:
            ind.set(qn("w:hanging"), str(hanging))
    if jc is not None:
        _child(ppr, "w:jc", val=jc)
    if outline is not None:
        _child(ppr, "w:outlineLvl", val=outline)


def _rpr_of_style(st):
    return _child(st, "w:rPr")


def _ppr_of_style(st):
    return _child(st, "w:pPr")


def _run(text: str, *, west: str = TNR, east: str = SONG, size: int | None = None,
         bold: bool = False, underline: bool = False):
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    r.append(rpr)
    _set_fonts(rpr, west, east)
    if bold:
        _set_bold(rpr, True)
    if size:
        _set_size(rpr, size)
    if underline:
        _child(rpr, "w:u", val="single")
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    r.append(t)
    return r


def _tab_run():
    r = OxmlElement("w:r")
    r.append(OxmlElement("w:tab"))
    return r


def _para(style: str | None = None):
    p = OxmlElement("w:p")
    if style:
        _child(_ppr(p), "w:pStyle", val=style)
    return p


def _p_text(p) -> str:
    return "".join(t.text or "" for t in p.iter(qn("w:t")))


def _p_style(p) -> str:
    ps = p.find(f"{qn('w:pPr')}/{qn('w:pStyle')}")
    return ps.get(qn("w:val")) if ps is not None else ""


def _page_break_para():
    p = OxmlElement("w:p")
    r = OxmlElement("w:r")
    br = OxmlElement("w:br")
    br.set(qn("w:type"), "page")
    r.append(br)
    p.append(r)
    return p


# --------------------------------------------------------------------------- 样式（官方正文格式）
def _ensure_style(d, style_id: str, name: str, kind: str = "paragraph", based_on: str | None = None):
    styles_el = d.styles.element
    for st in styles_el.findall(qn("w:style")):
        if st.get(qn("w:styleId")) == style_id:
            return st
        nm = st.find(qn("w:name"))
        if nm is not None and nm.get(qn("w:val")).lower() == name.lower():
            return st
    st = OxmlElement("w:style")
    st.set(qn("w:type"), kind)
    st.set(qn("w:styleId"), style_id)
    _child(st, "w:name", val=name)
    if based_on:
        _child(st, "w:basedOn", val=based_on)
    _child(st, "w:qFormat")
    styles_el.append(st)
    return st


def _set_pstyle(p, style_id: str) -> None:
    ppr = _ppr(p)
    _clear(ppr, "w:pStyle")
    _child(ppr, "w:pStyle", val=style_id)


def _find_style(d, *names: str):
    styles_el = d.styles.element
    lowered = {n.lower() for n in names}
    for st in styles_el.findall(qn("w:style")):
        if st.get(qn("w:styleId")) in names:
            return st
        nm = st.find(qn("w:name"))
        if nm is not None and nm.get(qn("w:val")).lower() in lowered:
            return st
    return None


def _style_text(st, *, west: str = TNR, east: str = SONG, size: int | None = None,
                bold: bool | None = None, color_black: bool = True, italic: bool | None = None) -> None:
    rpr = _rpr_of_style(st)
    _set_fonts(rpr, west, east)
    if size is not None:
        _set_size(rpr, size)
    if bold is not None:
        _set_bold(rpr, bold)
    if italic is not None:
        _clear(rpr, "w:i", "w:iCs")
        if italic:
            _child(rpr, "w:i")
    if color_black:
        _clear(rpr, "w:color", "w:shd")
    _child(rpr, "w:kern", val=2)
    _child(rpr, "w:lang", val="en-US", eastAsia="zh-CN")


def apply_body_styles(d) -> None:
    """把输出文档中 pandoc 用到的样式改成官方正文格式（不依赖 reference.docx 的名字对得上）。"""
    # 正文族：Normal / Body Text / First Paragraph（首行缩进 2 字符，两端对齐，单倍行距，段前后 0）
    for names in (("Normal",), ("Body Text", "BodyText"), ("First Paragraph", "FirstParagraph")):
        st = _find_style(d, *names)
        if st is None:
            continue
        _style_text(st, size=24, bold=False, italic=False)
        ppr = _ppr_of_style(st)
        _clear(ppr, "w:spacing", "w:ind", "w:jc")
        _para_props(ppr, jc="both", first_line_chars=200, before=0, after=0, line=240, line_rule="auto")
    # Compact（列表项 / 单元格）：无缩进
    st = _find_style(d, "Compact")
    if st is not None:
        _style_text(st, size=24)
        ppr = _ppr_of_style(st)
        _clear(ppr, "w:ind", "w:spacing")
        _para_props(ppr, first_line_chars=0, before=60, after=60)
    # 标题：一级 黑体 14 pt 居中；二级 12 pt 加粗；三级 12 pt
    spec = {
        1: dict(east=HEI, size=28, bold=False, jc="center"),
        2: dict(east=SONG, size=24, bold=True, jc="left"),
        3: dict(east=SONG, size=24, bold=False, jc="left"),
    }
    for lvl, sp in spec.items():
        st = _find_style(d, f"Heading {lvl}", f"Heading{lvl}", f"heading {lvl}")
        if st is None:
            continue
        _style_text(st, east=sp["east"], size=sp["size"], bold=sp["bold"], italic=False)
        ppr = _ppr_of_style(st)
        _clear(ppr, "w:ind", "w:spacing", "w:jc", "w:pBdr", "w:numPr")
        _para_props(ppr, jc=sp["jc"], first_line_chars=0, before=120, after=120, line=240, line_rule="auto",
                    keep_next=True, keep_lines=True, outline=lvl - 1)
    for lvl in range(4, 10):
        st = _find_style(d, f"Heading {lvl}", f"Heading{lvl}")
        if st is not None:
            _style_text(st, size=24, bold=True, italic=False)
    # 图题 / 表题：11 pt 加粗居中，无缩进
    for names, before, after in ((("Image Caption", "ImageCaption"), 60, 120),
                                 (("Table Caption", "TableCaption"), 120, 60),
                                 (("Caption",), 60, 120)):
        st = _find_style(d, *names)
        if st is None:
            continue
        _style_text(st, size=22, bold=True, italic=False)
        ppr = _ppr_of_style(st)
        _clear(ppr, "w:ind", "w:spacing", "w:jc")
        _para_props(ppr, jc="center", first_line_chars=0, before=before, after=after, line=240, line_rule="auto",
                    keep_next=names[0].startswith("Table"), keep_lines=True)
    # 图片段落：居中、无缩进、与图题同页
    for names in (("Captioned Figure", "CaptionedFigure"), ("Figure",)):
        st = _find_style(d, *names)
        if st is None:
            continue
        ppr = _ppr_of_style(st)
        _clear(ppr, "w:ind", "w:spacing", "w:jc")
        _para_props(ppr, jc="center", first_line_chars=0, before=120, after=60, keep_next=True)
    # 代码清单：等宽 10.5 pt，0.5 pt 全框线，单倍行距，无缩进
    st = _find_style(d, "Source Code", "SourceCode")
    if st is not None:
        _style_text(st, west=CODE_FONT, east=SONG, size=21, bold=False)
        ppr = _ppr_of_style(st)
        _clear(ppr, "w:ind", "w:spacing", "w:jc", "w:pBdr", "w:shd")
        _para_props(ppr, jc="left", first_line_chars=0, before=0, after=0, line=240, line_rule="auto")
        bdr = _child(ppr, "w:pBdr")
        for side in ("top", "left", "bottom", "right"):
            _child(bdr, f"w:{side}", val="single", sz=4, space=4, color="auto")
    st = _find_style(d, "Verbatim Char", "VerbatimChar")
    if st is not None:
        rpr = _rpr_of_style(st)
        _set_fonts(rpr, CODE_FONT, SONG)
        _set_size(rpr, 21)
    # 参考文献条目：12 pt 两端对齐，固定行距 18 pt，悬挂缩进 482
    st = _ensure_style(d, "References", "References", based_on="Normal")
    _style_text(st, size=24, bold=False)
    ppr = _ppr_of_style(st)
    _clear(ppr, "w:ind", "w:spacing", "w:jc")
    _para_props(ppr, jc="both", left=482, hanging=482, before=0, after=0, line=360, line_rule="exact")
    # 目录
    st = _find_style(d, "TOC Heading", "TOCHeading")
    if st is None:
        st = _ensure_style(d, "TOCHeading", "TOC Heading", based_on="Normal")
    _style_text(st, east=SONG, size=32, bold=True, italic=False)
    ppr = _ppr_of_style(st)
    _clear(ppr, "w:ind", "w:spacing", "w:jc", "w:numPr", "w:outlineLvl", "w:pBdr")
    _para_props(ppr, jc="center", first_line_chars=0, before=0, after=120, line=400, line_rule="exact")
    for lvl, left in ((1, 0), (2, 240), (3, 480)):
        st = _find_style(d, f"toc {lvl}", f"TOC{lvl}")
        if st is None:
            st = _ensure_style(d, f"TOC{lvl}", f"toc {lvl}", based_on="Normal")
        _style_text(st, size=21, bold=False, italic=False)
        ppr = _ppr_of_style(st)
        _clear(ppr, "w:ind", "w:spacing", "w:jc", "w:tabs")
        _para_props(ppr, jc="both", first_line_chars=0, left=left, before=60, after=60, line=240, line_rule="auto")
        tabs = _child(ppr, "w:tabs")
        _child(tabs, "w:tab", val="right", leader="dot", pos=9060)
    # 页眉 / 页脚：9 pt 居中
    for sid, name in (("Header", "header"), ("Footer", "footer")):
        st = _find_style(d, name, sid)
        if st is None:
            st = _ensure_style(d, sid, name, based_on="Normal")
        _style_text(st, size=18, bold=False)
        ppr = _ppr_of_style(st)
        _clear(ppr, "w:ind", "w:spacing", "w:jc", "w:tabs", "w:pBdr")
        _para_props(ppr, jc="center", first_line_chars=0, before=60, after=60, line=240, line_rule="auto")
    # 摘要页专用样式
    st = _ensure_style(d, "AbstractBody", "Abstract Body", based_on="Normal")
    _style_text(st, size=24, bold=False)
    ppr = _ppr_of_style(st)
    _clear(ppr, "w:ind", "w:spacing", "w:jc")
    _para_props(ppr, jc="both", first_line_chars=200, before=60, after=60, line=240, line_rule="auto")
    # 公式编号 / 表格正文
    st = _ensure_style(d, "TableText", "Table Text", based_on="Normal")
    _style_text(st, size=24, bold=False)
    ppr = _ppr_of_style(st)
    _clear(ppr, "w:ind", "w:spacing", "w:jc")
    _para_props(ppr, jc="center", first_line_chars=0, before=0, after=0, line=360, line_rule="exact")
    st = _ensure_style(d, "Equation", "Equation", based_on="Normal")
    _style_text(st, size=24, bold=False)
    ppr = _ppr_of_style(st)
    _clear(ppr, "w:ind", "w:spacing", "w:jc")
    _para_props(ppr, jc="center", first_line_chars=0, before=60, after=60, line=240, line_rule="auto")
    # 文档默认字体
    dd = d.styles.element.find(qn("w:docDefaults"))
    if dd is not None:
        rpd = dd.find(qn("w:rPrDefault"))
        if rpd is not None:
            rpr = _child(rpd, "w:rPr")
            _set_fonts(rpr, TNR, SONG)
            _set_size(rpr, 24)


def apply_page_setup(d, *, cover_placeholder: bool) -> None:
    """A4、四边 2.5 cm、页眉/页脚距、页脚居中 9 pt 页码；可选空白封面节（页码从摘要页起为 1）。"""
    body = d.element.body
    sect = body.find(qn("w:sectPr"))
    if sect is None:
        sect = OxmlElement("w:sectPr")
        body.append(sect)
    s = d.sections[-1]
    s.page_width, s.page_height = Twips(PAGE_W), Twips(PAGE_H)
    s.top_margin = s.bottom_margin = s.left_margin = s.right_margin = Twips(MARGIN)
    s.header_distance, s.footer_distance = Twips(HEADER_DIST), Twips(FOOTER_DIST)
    s.footer.is_linked_to_previous = False
    s.header.is_linked_to_previous = False
    hp = s.header.paragraphs[0] if s.header.paragraphs else s.header.add_paragraph()
    for r in list(hp._p):
        if r.tag != qn("w:pPr"):
            hp._p.remove(r)
    _set_pstyle(hp._p, "Header")
    fp = s.footer.paragraphs[0] if s.footer.paragraphs else s.footer.add_paragraph()
    for r in list(fp._p):
        if r.tag != qn("w:pPr"):
            fp._p.remove(r)
    _set_pstyle(fp._p, "Footer")
    _para_props(_ppr(fp._p), jc="center")
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), " PAGE \\* MERGEFORMAT ")
    fld.append(_run("1", size=18))
    fp._p.append(fld)
    _clear(sect, "w:pgNumType", "w:titlePg")
    _child(sect, "w:pgNumType", start=1)

    if cover_placeholder:
        cover = _para()
        _para_props(_ppr(cover), jc="center", first_line_chars=0, before=4000)
        cover.append(_run("[封面：请用当届官方模板封面替换本页，删除本行]", east=HEI, size=28, bold=True))
        brk = _para()
        csect = copy.deepcopy(sect)
        _clear(csect, "w:headerReference", "w:footerReference", "w:pgNumType")
        _child(csect, "w:pgNumType", start=0)
        _ppr(brk).append(csect)
        body.insert(0, brk)
        body.insert(0, cover)


# --------------------------------------------------------------------------- 正文结构后处理
def _iter_body_paras(body):
    return [p for p in body.iterchildren(qn("w:p"))]


def fix_heading_numbers(body) -> list[tuple[int, str]]:
    """pandoc --number-sections 输出 `3<tab>标题` → 一级改为 `三、 标题`，二/三级 `1.1 标题`。

    返回 [(level, 带编号标题文本)]，供目录缓存条目使用。
    """
    headings: list[tuple[int, str]] = []
    for p in body.iter(qn("w:p")):
        st = _p_style(p)
        m = re.fullmatch(r"Heading(\d)", st)
        if not m:
            continue
        lvl = int(m.group(1))
        runs = [r for r in p.findall(qn("w:r"))]
        num_run = next((r for r in runs if r.find(f"{qn('w:rPr')}/{qn('w:rStyle')}") is not None
                        and r.find(f"{qn('w:rPr')}/{qn('w:rStyle')}").get(qn("w:val")) == "SectionNumber"), None)
        if num_run is not None:
            t = num_run.find(qn("w:t"))
            if lvl == 1 and t is not None and t.text and t.text.strip().isdigit():
                t.text = cn_number(int(t.text.strip())) + "、"
            rpr = num_run.find(qn("w:rPr"))
            _set_fonts(rpr, TNR, HEI if lvl == 1 else SONG)
            idx = list(p).index(num_run)
            nxt = p[idx + 1] if idx + 1 < len(p) else None
            if nxt is not None and nxt.tag == qn("w:r") and nxt.find(qn("w:tab")) is not None:
                p.remove(nxt)
                if lvl > 1:
                    t.text = (t.text or "") + " "
        if lvl <= 3:
            headings.append((lvl, _p_text(p).strip()))
    return headings


def style_front_matter(body, fm: FrontMatter) -> None:
    """按官方摘要页格式重排 标题行 / “摘 要：” / 摘要段 / 关键词行，并删除零宽标记。"""
    paras = _iter_body_paras(body)
    title_p = abs_p = kw_p = None
    for p in paras[:12]:
        txt = _p_text(p)
        if TITLE_MARK in txt:
            title_p = p
        elif ABSTRACT_MARK in txt:
            abs_p = p
        elif KEYWORDS_MARK in txt:
            kw_p = p
    if title_p is not None:
        title = _p_text(title_p).replace(TITLE_MARK, "").strip()
        for el in list(title_p):
            if el.tag != qn("w:pPr"):
                title_p.remove(el)
        ppr = _ppr(title_p)
        _clear(ppr, "w:ind", "w:jc", "w:spacing")
        _set_pstyle(title_p, "Normal")
        _para_props(ppr, jc="left", first_line_chars=0, before=120, after=120, line=360, line_rule="auto")
        title_p.append(_run("题\u3000目：", east=LI, size=36))
        title_p.append(_run(title, size=28, underline=True))
    if abs_p is not None:
        for el in list(abs_p):
            if el.tag != qn("w:pPr"):
                abs_p.remove(el)
        ppr = _ppr(abs_p)
        _clear(ppr, "w:ind", "w:jc", "w:spacing")
        _set_pstyle(abs_p, "Normal")
        _para_props(ppr, jc="center", first_line_chars=0, before=120, after=120, line=360, line_rule="auto")
        abs_p.append(_run("摘\u3000要：", east=LI, size=36))
    if abs_p is not None and kw_p is not None:
        start, end = paras.index(abs_p), paras.index(kw_p)
        for p in paras[start + 1:end]:
            ppr = _ppr(p)
            _clear(ppr, "w:ind", "w:spacing")
            _set_pstyle(p, "AbstractBody")
            # “针对问题一，”引导词加粗
            first_r = p.find(qn("w:r"))
            first_t = first_r.find(qn("w:t")) if first_r is not None else None
            if first_t is not None and first_t.text:
                m = LEAD_PHRASE.match(first_t.text)
                if m:
                    first_t.text = first_t.text[m.end():]
                    first_r.addprevious(_run(m.group(1), size=24, bold=True))
    if kw_p is not None:
        keywords = fm.keywords or ["[关键词]"]
        for el in list(kw_p):
            if el.tag != qn("w:pPr"):
                kw_p.remove(el)
        ppr = _ppr(kw_p)
        _clear(ppr, "w:ind", "w:jc", "w:spacing")
        _set_pstyle(kw_p, "Normal")
        _para_props(ppr, jc="left", first_line_chars=0, before=240, after=120, line=360, line_rule="auto")
        kw_p.append(_run("关键词：", east=LI, size=36))
        kw_p.append(_run("\u3000\u3000".join(keywords), size=24))


def insert_toc(body, headings: list[tuple[int, str]], pages: dict[str, int] | None = None) -> None:
    """在第一个一级标题前插入目录页：`目录` 标题 + TOC 域（带静态缓存条目）+ 分页。"""
    first_h1 = next((p for p in body.iter(qn("w:p")) if _p_style(p) == "Heading1"), None)
    if first_h1 is None or not headings:
        return
    anchor = first_h1
    title_p = _para("TOCHeading")
    title_p.append(_run("目录", size=32, bold=True))
    anchor.addprevious(title_p)
    n = len(headings)
    for i, (lvl, text) in enumerate(headings):
        p = _para(f"TOC{lvl}")
        if i == 0:
            r = OxmlElement("w:r")
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), "begin")
            fc.set(qn("w:dirty"), "true")
            r.append(fc)
            p.append(r)
            r = OxmlElement("w:r")
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = ' TOC \\o "1-3" \\h \\z \\u '
            r.append(it)
            p.append(r)
            r = OxmlElement("w:r")
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), "separate")
            r.append(fc)
            p.append(r)
        p.append(_run(text, size=21))
        p.append(_tab_run())
        pg = (pages or {}).get(text)
        p.append(_run(str(pg) if pg else "", size=21))
        if i == n - 1:
            r = OxmlElement("w:r")
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), "end")
            r.append(fc)
            p.append(r)
        anchor.addprevious(p)
    anchor.addprevious(_page_break_para())


EQ_TABLE_MARK = "hwb-equation"


def _is_eq_table(tbl) -> bool:
    cap = tbl.find(f"{qn('w:tblPr')}/{qn('w:tblCaption')}")
    return cap is not None and cap.get(qn("w:val")) == EQ_TABLE_MARK


def _equation_table(omath_para, number: int):
    tbl = OxmlElement("w:tbl")
    tblpr = _child(tbl, "w:tblPr")
    _child(tblpr, "w:tblW", w=5000, type="pct")
    _child(tblpr, "w:jc", val="center")
    bd = _child(tblpr, "w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        _child(bd, f"w:{side}", val="nil")
    _child(tblpr, "w:tblLayout", type="fixed")
    mar = _child(tblpr, "w:tblCellMar")
    for side in ("left", "right"):
        _child(mar, f"w:{side}", w=0, type="dxa")
    _child(tblpr, "w:tblLook", val="0000", firstRow=0, lastRow=0, firstColumn=0, lastColumn=0, noHBand=1, noVBand=1)
    _child(tblpr, "w:tblCaption", val=EQ_TABLE_MARK)
    grid = OxmlElement("w:tblGrid")
    tbl.append(grid)
    side_w = TEXT_W // 10
    widths = (side_w, TEXT_W - 2 * side_w, side_w)
    for w in widths:
        _child(grid, "w:gridCol", w=w)
    tr = OxmlElement("w:tr")
    tbl.append(tr)
    for i, w in enumerate(widths):
        tc = OxmlElement("w:tc")
        tcpr = _child(tc, "w:tcPr")
        _child(tcpr, "w:tcW", w=w, type="dxa")
        _child(tcpr, "w:vAlign", val="center")
        p = _para("Equation")
        if i == 1:
            p.append(omath_para)
        elif i == 2:
            _para_props(_ppr(p), jc="right")
            p.append(_run(f"({number})", size=24))
        tc.append(p)
        tr.append(tc)
    return tbl


def number_equations(body) -> int:
    """行间公式（oMathPara）→ 三栏无框表：空 | 公式居中 | (N) 右对齐。返回公式数。"""
    n = 0
    made: list = []
    for omp in list(body.iter(qn("m:oMathPara"))):
        p = omp.getparent()
        if p is None or p.tag != qn("w:p"):
            continue
        # 表格内的公式保持原样
        anc = p.getparent()
        if anc is not None and anc.tag == qn("w:tc"):
            continue
        n += 1
        tbl = _equation_table(omp, n)
        p.addprevious(tbl)
        made.append(tbl)
        if not _p_text(p).strip() and p.find(qn("m:oMathPara")) is None and p.find(qn("w:drawing")) is None:
            p.getparent().remove(p)
    # 相邻两个表在 Word 中会合并，中间垫一个极小空段
    for tbl in made:
        for sib in (tbl.getnext(), tbl.getprevious()):
            if sib is not None and sib.tag == qn("w:tbl"):
                spacer = _para()
                _para_props(_ppr(spacer), before=0, after=0, line=20, line_rule="exact")
                _child(_child(_ppr(spacer), "w:rPr"), "w:sz", val=2)
                tbl.addnext(spacer) if sib is tbl.getnext() else tbl.addprevious(spacer)
    return n


def style_tables(body) -> int:
    """普通表 → 三线表：居中，顶线/底线 1.5 pt，表头下线 0.5 pt，表头加粗，单元格 12 pt 固定行距 18 pt。"""
    n = 0
    for tbl in body.iter(qn("w:tbl")):
        tblpr = tbl.find(qn("w:tblPr"))
        if tblpr is None or _is_eq_table(tbl):
            continue
        if tbl.getparent().tag == qn("w:tc"):
            continue  # 嵌套表
        n += 1
        _child(tblpr, "w:jc", val="center")
        _clear(tblpr, "w:tblBorders")
        bd = _child(tblpr, "w:tblBorders")
        for side, val, sz in (("top", "single", 12), ("left", "nil", 0), ("bottom", "single", 12),
                              ("right", "nil", 0), ("insideH", "nil", 0), ("insideV", "nil", 0)):
            el = _child(bd, f"w:{side}", val=val)
            if val != "nil":
                el.set(qn("w:sz"), str(sz))
                el.set(qn("w:space"), "0")
                el.set(qn("w:color"), "auto")
        rows = tbl.findall(qn("w:tr"))
        if not rows:
            continue
        header_rows = [r for r in rows if r.find(f"{qn('w:trPr')}/{qn('w:tblHeader')}") is not None] or rows[:1]
        last_header = header_rows[-1]
        for tr in rows:
            is_header = tr in header_rows
            for tc in tr.findall(qn("w:tc")):
                tcpr = _child(tc, "w:tcPr")
                if tr is last_header:
                    tb = _child(tcpr, "w:tcBorders")
                    _child(tb, "w:bottom", val="single", sz=4, space=0, color="auto")
                _child(tcpr, "w:vAlign", val="center")
                for p in tc.findall(qn("w:p")):
                    ppr = _ppr(p)
                    jc = ppr.find(qn("w:jc"))
                    jc_val = jc.get(qn("w:val")) if jc is not None else "center"
                    _clear(ppr, "w:ind", "w:spacing", "w:jc")
                    _set_pstyle(p, "TableText")
                    _para_props(ppr, jc=jc_val, first_line_chars=0, before=0, after=0, line=360, line_rule="exact")
                    if is_header:
                        for r in p.findall(qn("w:r")):
                            _set_bold(_child(r, "w:rPr"), True)
    return n


def style_references(body) -> int:
    """`参考文献` 一级标题到下一个一级标题之间的段落套 References 样式。返回条目数。"""
    n = 0
    in_refs = False
    for p in _iter_body_paras(body):
        st = _p_style(p)
        if st == "Heading1":
            in_refs = _p_text(p).strip().startswith("参考文献")
            continue
        if in_refs and st in ("", "Normal", "BodyText", "FirstParagraph", "Compact") and _p_text(p).strip():
            _clear(_ppr(p), "w:ind", "w:spacing")
            _set_pstyle(p, "References")
            n += 1
    return n


def apply_body_format(docx_path: Path, fm: FrontMatter, *, front_matter: bool, cover_placeholder: bool,
                         page_setup: bool, toc_pages: dict[str, int] | None = None) -> dict:
    d = docx.Document(str(docx_path))
    body = d.element.body
    apply_body_styles(d)
    headings = fix_heading_numbers(body)
    if front_matter:
        style_front_matter(body, fm)
    n_eq = number_equations(body)
    n_tbl = style_tables(body)
    n_ref = style_references(body)
    insert_toc(body, headings, toc_pages)
    if page_setup:
        apply_page_setup(d, cover_placeholder=cover_placeholder)
    d.save(str(docx_path))
    return {"numbered_equations": n_eq, "three_line_tables": n_tbl, "reference_entries": n_ref,
            "toc_entries": [t for _, t in headings]}


# --------------------------------------------------------------------------- reference.docx
def build_reference_docx(pandoc: str, dst: Path) -> None:
    """pandoc 默认 reference.docx + A4 页面（pandoc 据此缩放图片）；样式细节在输出文档上统一施加。"""
    raw = run([pandoc, "-o", str(dst), "--print-default-data-file", "reference.docx"])
    if raw.returncode != 0:
        sys.exit(f"pandoc 生成 reference.docx 失败: {raw.stderr}")
    d = docx.Document(str(dst))
    apply_body_styles(d)
    s = d.sections[0]
    s.page_width, s.page_height = Twips(PAGE_W), Twips(PAGE_H)
    s.top_margin = s.bottom_margin = s.left_margin = s.right_margin = Twips(MARGIN)
    s.header_distance, s.footer_distance = Twips(HEADER_DIST), Twips(FOOTER_DIST)
    d.save(str(dst))


# --------------------------------------------------------------------------- 检查
def inspect(docx_path: Path) -> dict:
    z = zipfile.ZipFile(docx_path)
    xml = z.read("word/document.xml").decode("utf-8")
    text = re.sub(r"<[^>]+>", "", xml)
    footers = "".join(z.read(n).decode("utf-8") for n in z.namelist() if re.match(r"word/footer\d*\.xml", n))
    sect = re.findall(r"<w:sectPr.*?</w:sectPr>", xml, flags=re.S)
    last_sect = sect[-1] if sect else ""
    pgsz = re.search(r'<w:pgSz[^>]*w:w="(\d+)"[^>]*w:h="(\d+)"', last_sect) or re.search(r'<w:pgSz[^>]*w:h="(\d+)"[^>]*w:w="(\d+)"', last_sect)
    pgmar = re.search(r"<w:pgMar([^>]*)/>", last_sect)
    margins = dict(re.findall(r'w:(top|bottom|left|right)="(\d+)"', pgmar.group(1))) if pgmar else {}
    styles_used = re.findall(r'<w:pStyle w:val="([^"]+)"', xml)
    fig_paras = styles_used.count("CaptionedFigure") + styles_used.count("Figure")
    n_eq_tables = xml.count(f'<w:tblCaption w:val="{EQ_TABLE_MARK}"')
    return {
        "paragraphs": xml.count("<w:p>") + xml.count("<w:p "),
        "images": sum(1 for n in z.namelist() if n.startswith("word/media/")),
        "equations": xml.count("<m:oMath>") + xml.count("<m:oMath "),
        "display_equations": xml.count("<m:oMathPara>") + xml.count("<m:oMathPara "),
        "tables": xml.count("<w:tbl>") - n_eq_tables,
        "headings": {f"h{l}": styles_used.count(f"Heading{l}") for l in (1, 2, 3)},
        "image_captions": styles_used.count("ImageCaption"),
        "table_captions": styles_used.count("TableCaption"),
        "uncaptioned_figures": styles_used.count("Figure"),
        "figure_paragraphs": fig_paras,
        "page_size_twips": [int(pgsz.group(1)), int(pgsz.group(2))] if pgsz else None,
        "margins_twips": {k: int(v) for k, v in margins.items()},
        "footer_page_field": "PAGE" in footers,
        "toc_field": "TOC \\o" in xml,
        "placeholders": sorted(set(PLACEHOLDER.findall(text))),
        "leaked_internal_names": sorted({w for w in INTERNAL_NAMES if w in text}),
    }


def audit(report: dict, *, expect_page_setup: bool) -> tuple[list[str], list[str]]:
    """返回 (errors, warnings)：errors 在 --strict 下非零退出，warnings 仅提示。"""
    errors: list[str] = []
    warnings: list[str] = []
    if report["placeholders"]:
        errors.append("DOCX 中仍有模板占位符，提交前必须替换: " + "、".join(report["placeholders"]))
    if report["leaked_internal_names"]:
        errors.append("DOCX 中出现工作流内部文件名，提交前必须删除: " + "、".join(report["leaked_internal_names"]))
    if expect_page_setup:
        if sorted(report["page_size_twips"] or []) != sorted([PAGE_W, PAGE_H]):
            errors.append(f"页面尺寸不是 A4: {report['page_size_twips']}")
        if len(report["margins_twips"]) < 4 or any(v != MARGIN for v in report["margins_twips"].values()):
            errors.append(f"页边距不是四边 2.5 cm: {report['margins_twips']}")
        if not report["footer_page_field"]:
            errors.append("页脚缺少 PAGE 页码域")
    if report["headings"]["h1"] == 0:
        errors.append("没有一级标题（正文未按章节组织？）")
    if report["uncaptioned_figures"]:
        warnings.append(f"{report['uncaptioned_figures']} 张图没有图题")
    if report["images"] and report["image_captions"] < report["images"] - report["uncaptioned_figures"]:
        warnings.append("图题数少于图片数（多图并列时可忽略）")
    if report["tables"] and report["table_captions"] < report["tables"]:
        warnings.append(f"表题数({report['table_captions']})少于表格数({report['tables']})")
    if not report["toc_field"]:
        warnings.append("未生成目录")
    return errors, warnings


def docx_to_pdf(docx_path: Path) -> Path | None:
    soffice = find_soffice()
    if not soffice:
        print("[info] 未找到 soffice，跳过 DOCX->PDF 渲染抽检")
        return None
    # 输出到独立子目录，避免覆盖排版引擎生成的 main.pdf
    outdir = docx_path.parent / "docx_render"
    outdir.mkdir(exist_ok=True)
    r = run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(outdir), str(docx_path)], timeout=300)
    pdf = outdir / docx_path.with_suffix(".pdf").name
    if r.returncode != 0 or not pdf.exists():
        print(f"[warn] soffice 转换失败: {r.stderr[-500:]}")
        return None
    return pdf


def locate_heading_pages(pdf: Path, headings: list[str], page_offset: int) -> dict[str, int]:
    """在渲染 PDF 里按行精确匹配标题文本，返回 {标题: 显示页码}（跳过目录页的点线条目）。"""
    try:
        import pymupdf  # type: ignore
    except ImportError:
        return {}
    doc = pymupdf.open(pdf)
    norm = lambda s: re.sub(r"\s+", "", s)
    targets = [norm(h) for h in headings]
    found: dict[str, int] = {}
    start_page = 0
    for h, key in zip(headings, targets):
        for pno in range(start_page, doc.page_count):
            lines = [norm(l) for l in doc[pno].get_text().splitlines()]
            if key in lines:
                found[h] = pno + 1 - page_offset
                start_page = pno
                break
    doc.close()
    return found



# --------------------------------------------------------------------------- 主流程
def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--paper", default="paper", help="论文目录（含 paper.yaml 与 sections/*.md）")
    ap.add_argument("--out", default=None, help="输出 DOCX 路径，默认 <paper>/main.docx")
    ap.add_argument("--reference", default=None, help="Word 参考样式文档（当届官方 DOCX 模板可直接传入；页面设置以其为准）")
    ap.add_argument("--title", default=None)
    ap.add_argument("--keywords", default=None, help="分号分隔，覆盖 paper.yaml")
    ap.add_argument("--no-front-matter", action="store_true", help="不生成题目/摘要/关键词页（例如官方模板已含）")
    ap.add_argument("--cover-placeholder", action="store_true", help="最前面加一页空白封面节（页码 0，不显示），摘要页从 1 起")
    ap.add_argument("--figure-dpi", type=int, default=300, help="PDF 图转 PNG 的分辩率（默认 300）")
    ap.add_argument("--pdf", action="store_true", help="用 soffice 顺带导出 PDF 做渲染抽检，并回填目录页码")
    ap.add_argument("--keep-entry", action="store_true", help="pandoc 失败时保留 _docx_body.md 中间文件")
    ap.add_argument("--strict", action="store_true", help="发现占位符、内部文件名泄露或格式审计问题时以非零退出（终稿检查用）")
    ap.add_argument("--freeze", action="store_true", help="生成后冻结：Word 成为唯一真源")
    ap.add_argument("--unfreeze", action="store_true", help="解除冻结后退出（不生成）")
    ap.add_argument("--status", action="store_true", help="只打印冻结状态")
    ap.add_argument("--force", action="store_true", help="已冻结仍重生成（覆盖人工 Word 修改，先备份为 main.docx.bak-<时间>）")
    args = ap.parse_args()

    paper = Path(args.paper).resolve()
    out = Path(args.out).resolve() if args.out else paper / "main.docx"

    if args.status:
        print_status(paper, out)
        return
    if args.unfreeze:
        rec = load_freeze(paper)
        rec.update({"frozen": False, "unfrozen_at": _dt.datetime.now().isoformat(timespec="seconds")})
        save_freeze(paper, rec)
        print(f"[info] 已解除冻结；注意 Word 中的人工修改不会回流到 Markdown，重生成前请先把改动同步到 sections/*.md")
        return

    st, rec = freeze_status(paper, out)
    if st.startswith("frozen"):
        print_status(paper, out)
        if not args.force:
            sys.exit("[error] DOCX 已冻结：请直接编辑 Word；确需重生成加 --force（会备份并覆盖），或 --unfreeze。")
        if out.exists():
            bak = out.with_name(out.name + ".bak-" + _dt.datetime.now().strftime("%Y%m%d-%H%M%S"))
            shutil.copy2(out, bak)
            print(f"[info] 已备份被覆盖的 Word: {bak}")

    pandoc = find_pandoc()
    if not pandoc:
        sys.exit("缺少 pandoc：Windows 到 https://pandoc.org/installing.html 安装（或 winget install JohnMacFarlane.Pandoc）；"
                 "Linux 运行 bash scripts/setup_env.sh；也可设 PANDOC_BIN")

    meta = load_meta(paper)
    abstract_file, sections = collect_sections(paper, meta)
    if not sections:
        sys.exit("paper/sections/ 下没有正文章节 .md")
    print(f"[info] 章节 {len(sections)} 个" + (f"，摘要 {abstract_file.name}" if abstract_file else "，无摘要文件"))

    fm = extract_front_matter(paper, meta, abstract_file)
    if args.title:
        fm.title = args.title
    if args.keywords:
        fm.keywords = [k.strip() for k in re.split(r"[;；]", args.keywords) if k.strip()]
    print(f"[info] 标题={fm.title!r} 摘要段数={len(fm.abstract)} 关键词={fm.keywords}")

    all_files = ([abstract_file] if abstract_file else []) + sections
    numbers, texts = resolve_crossrefs(sections)
    if numbers:
        print(f"[info] 交叉引用 {len(numbers)} 个：" + "、".join(f"{k}→{v}" for k, v in list(numbers.items())[:8]) + ("…" if len(numbers) > 8 else ""))
    sec_dir = paper / "sections"
    pngs = convert_pdf_figures(texts, sec_dir, dpi=args.figure_dpi)
    if pngs:
        print(f"[info] 生成 PNG 图 {len(pngs)} 张")

    ref_arg = args.reference or meta.get("reference")
    with tempfile.TemporaryDirectory() as td:
        tdp = Path(td)
        lua = tdp / "draft_docx.lua"
        lua.write_text(PDF2PNG_LUA, encoding="utf-8")
        ref = (paper / str(ref_arg)).resolve() if ref_arg and not Path(str(ref_arg)).is_absolute() else (Path(str(ref_arg)) if ref_arg else tdp / "reference.docx")
        if not ref_arg:
            build_reference_docx(pandoc, ref)
        elif not ref.exists():
            sys.exit(f"参考模板不存在: {ref}")

        head = "" if args.no_front_matter else front_matter_source(fm)
        entry = sec_dir / "_docx_body.md"
        entry.write_text(head + "\n\n".join(texts) + "\n", encoding="utf-8")
        cmd = [
            pandoc, str(entry), "-f", "markdown+tex_math_dollars+raw_tex+implicit_figures+table_captions+pipe_tables+grid_tables+fenced_divs",
            "-t", "docx", "--number-sections",
            "--reference-doc", str(ref), "--lua-filter", str(lua),
            "--resource-path", os.pathsep.join(str(p) for p in (sec_dir, paper, paper.parent)),
            "-o", str(out),
        ]
        try:
            r = run(cmd, cwd=str(sec_dir))
        finally:
            if not args.keep_entry:
                entry.unlink(missing_ok=True)
        if r.returncode != 0:
            hint = f"（已保留 {entry} 供排查）" if args.keep_entry else "（加 --keep-entry 可保留中间文件排查）"
            sys.exit(f"pandoc 失败{hint}:\n{r.stderr}")
        if r.stderr.strip():
            print("[pandoc]", r.stderr.strip()[-800:])

    raw_docx = out.with_name(out.stem + "_pandoc_raw.docx")
    shutil.copy(out, raw_docx)
    page_setup = not ref_arg
    body_report = apply_body_format(out, fm, front_matter=not args.no_front_matter,
                                    cover_placeholder=args.cover_placeholder, page_setup=page_setup)

    report = inspect(out)
    report.update({k: v for k, v in body_report.items() if k != "toc_entries"})
    report["docx"] = str(out)
    if args.pdf:
        pdf = docx_to_pdf(out)
        if pdf:
            pages = locate_heading_pages(pdf, body_report["toc_entries"], 1 if args.cover_placeholder else 0)
            if pages:
                shutil.copy(raw_docx, out)
                apply_body_format(out, fm, front_matter=not args.no_front_matter,
                                  cover_placeholder=args.cover_placeholder, page_setup=page_setup, toc_pages=pages)
                pdf = docx_to_pdf(out) or pdf
                report["toc_pages_filled"] = len(pages)
            report["pdf"] = str(pdf)
            try:
                import pymupdf  # type: ignore

                report["pdf_pages"] = pymupdf.open(pdf).page_count
            except Exception:
                pass
    raw_docx.unlink(missing_ok=True)
    print(json.dumps(report, ensure_ascii=False, indent=2))
    errors, warnings = audit(report, expect_page_setup=page_setup)
    for w in warnings:
        print(f"[warn] {w}")
    for e in errors:
        print(f"[error] {e}")

    now = _dt.datetime.now().isoformat(timespec="seconds")
    rec = {
        "docx": out.name,
        "frozen": bool(args.freeze),
        "generated_at": now,
        "generated_by": "build_docx.py",
        "sections": [f.name for f in all_files],
        "sections_digest": sections_digest(all_files),
        "docx_sha256": sha256_of(out),
        "crossrefs": numbers,
        "audit_errors": errors,
        "audit_warnings": warnings,
        "rule": "frozen=true 时 Word 为唯一真源：只改 main.docx，不要重跑 build_docx.py（--force 会覆盖人工修改）",
    }
    if args.freeze:
        rec["frozen_at"] = now
        print(f"[info] 已冻结：从现在起只编辑 {out.name}，Markdown 仅作历史生成源")
    save_freeze(paper, rec)
    if errors and args.strict:
        sys.exit(2)


if __name__ == "__main__":
    main()
